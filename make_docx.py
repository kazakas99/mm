# ataskaita.tex -> ataskaita.docx (Word lygtys lieka redaguojamos)
import pypandoc
import docx
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH

s = open('ataskaita.tex', encoding='utf-8').read()
s = s.replace(r'\lstinputlisting{lab2_var9.py}', 'CODEPLACEHOLDER').replace(r'\newpage', '')
open('_word.tex', 'w', encoding='utf-8').write(s)
pypandoc.convert_file('_word.tex', 'docx', outputfile='ataskaita.docx',
                      extra_args=['--resource-path=.', '--number-sections'])

d = docx.Document('ataskaita.docx')
code = open('lab2_var9.py', encoding='utf-8').read().splitlines()
nf = nt = 0
for p in d.paragraphs:
    if 'CODEPLACEHOLDER' in p.text:
        p.text = ''
        for i, line in enumerate(code):
            r = p.add_run(line)
            r.font.name = 'Consolas'
            r.font.size = Pt(8.5)
            if i < len(code) - 1:
                r.add_break()
    if p.style.name in ('Image Caption', 'Caption'):
        nf += 1
        p.runs[0].text = f'{nf} pav. ' + p.runs[0].text
    if p.style.name == 'Table Caption':
        nt += 1
        p.runs[0].text = f'{nt} lentelė. ' + p.runs[0].text
    if p._p.xpath('.//pic:pic'):
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
for t in d.tables:
    t.alignment = 1  # centruota
d.save('ataskaita.docx')
print('docx ok; paveikslai:', nf, 'lentelės:', nt)
